"""File connectors: CSV and TSV, JSON and JSON Lines, Parquet (all read by DuckDB) and Excel workbooks (openpyxl)."""

from __future__ import annotations

import csv
import datetime as _dt
import os
import tempfile
from pathlib import Path
from typing import Dict, List, Optional

from ..schema import quote_ident
from ..workspace import Workspace
from . import Connector, ConnectorField, SourceError, SourceSpec, register

__all__ = ["stage_csv"]


def _path(spec: SourceSpec) -> Path:
    path = Path(str(spec.options.get("path", ""))).expanduser()
    if not path.is_file():
        raise SourceError(f"file not found: {path}")
    return path


def _types_literal(types: Optional[Dict[str, str]]) -> str:
    if not types:
        return ""
    items = ", ".join("'" + name.replace("'", "''") + "': '" + kind + "'" for name, kind in types.items())
    return f", types = {{{items}}}"


def stage_csv(workspace: Workspace, name: str, path: Path, limit: int, delimiter: str = "",
              types: Optional[Dict[str, str]] = None, header: bool = True) -> None:
    """Load a delimited file into table ``name``; DuckDB detects types from the whole file unless ``types`` pins some."""
    delim = f", delim = '{delimiter}'" if delimiter in (",", ";", "\t", "|") else ""
    workspace.execute(
        f"CREATE OR REPLACE TABLE {quote_ident(name)} AS SELECT * FROM read_csv(?, header = {str(header).lower()}, "
        f"sample_size = -1{delim}{_types_literal(types)}) LIMIT {int(limit)}", [str(path)])


def _stage_delimited(workspace: Workspace, spec: SourceSpec, _secret: str, name: str) -> None:
    path = _path(spec)
    delimiter = spec.options.get("delimiter") or ("\t" if path.suffix.lower() in (".tsv", ".tab") else "")
    stage_csv(workspace, name, path, spec.row_limit, delimiter, header=spec.options.get("header", True) is not False)


def _stage_json(workspace: Workspace, spec: SourceSpec, _secret: str, name: str) -> None:
    path = _path(spec)
    workspace.execute(f"CREATE OR REPLACE TABLE {quote_ident(name)} AS SELECT * FROM read_json_auto(?) "
                      f"LIMIT {int(spec.row_limit)}", [str(path)])


def _stage_parquet(workspace: Workspace, spec: SourceSpec, _secret: str, name: str) -> None:
    path = _path(spec)
    workspace.execute(f"CREATE OR REPLACE TABLE {quote_ident(name)} AS SELECT * FROM read_parquet(?) "
                      f"LIMIT {int(spec.row_limit)}", [str(path)])


def _excel_sheets(spec: SourceSpec, _secret: str = "") -> List[str]:
    import openpyxl
    workbook = openpyxl.load_workbook(_path(spec), read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _cell(value: object) -> object:
    if isinstance(value, _dt.datetime) and value.time() == _dt.time(0):
        return value.date().isoformat()
    if isinstance(value, (_dt.date, _dt.datetime)):
        return value.isoformat(sep=" ") if isinstance(value, _dt.datetime) else value.isoformat()
    return value


def _stage_excel(workspace: Workspace, spec: SourceSpec, _secret: str, name: str) -> None:
    import openpyxl
    path = _path(spec)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    handle, staging = tempfile.mkstemp(suffix=".csv", prefix="fds_")
    try:
        sheet_name = spec.options.get("sheet") or workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise SourceError(f"the workbook has no sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        written = 0
        with os.fdopen(handle, "w", newline="", encoding="utf-8") as out:
            writer = csv.writer(out)
            header: Optional[List[str]] = None
            for row in sheet.iter_rows(values_only=True):
                if header is None:
                    if not any(v is not None for v in row):
                        continue  # leading blank rows
                    header = [str(v) if v is not None else f"column_{i + 1}" for i, v in enumerate(row)]
                    writer.writerow(header)
                    continue
                if not any(v is not None for v in row):
                    continue
                writer.writerow([_cell(v) for v in row[:len(header)]])
                written += 1
                if written >= spec.row_limit:
                    break
        if header is None:
            raise SourceError(f"sheet {sheet_name!r} is empty")
        stage_csv(workspace, name, Path(staging), spec.row_limit, ",")
    finally:
        workbook.close()
        try:
            os.unlink(staging)
        except OSError:
            pass


_PATH = ConnectorField("path", "File", "path")

register(Connector("csv", "CSV or TSV file", "File",
                   [_PATH, ConnectorField("delimiter", "Delimiter", "choice", "", False, ["", ",", ";", "\t", "|"],
                                          "detect")],
                   _stage_delimited, extensions=[".csv", ".tsv", ".tab", ".txt"],
                   description="Delimited text; types are detected from the whole file."))
register(Connector("json", "JSON or JSON Lines file", "File", [_PATH], _stage_json,
                   extensions=[".json", ".jsonl", ".ndjson"],
                   description="An array of records or one record per line; nested fields become structs."))
register(Connector("parquet", "Parquet file", "File", [_PATH], _stage_parquet, extensions=[".parquet", ".pq"],
                   description="Columnar files with their own types."))
register(Connector("excel", "Excel workbook", "File",
                   [_PATH, ConnectorField("sheet", "Sheet", "choice", "", False)],
                   _stage_excel, tables=_excel_sheets, extensions=[".xlsx", ".xlsm"], requires="openpyxl",
                   description="One sheet per dataset; the first non-empty row is the header."))
